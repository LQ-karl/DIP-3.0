#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""fill_icd_names 补全逻辑的轻量回归测试（不依赖 data/ 大表）。"""
import os
import sys

import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import fill_icd_names as F  # noqa: E402


def test_lookup_exact_and_strip_x():
    mp = {"I50.900": "心力衰竭", "T25.300X003": "足部三度烧伤", "T31.000": "累及体表10%以下的烧伤"}
    assert F.lookup("I50.900", mp) == "心力衰竭"
    # 去掉 x 扩展码后命中
    assert F.lookup("T25.300x003", mp) == "足部三度烧伤"
    # 未命中返回空串
    assert F.lookup("Z99.999", mp) == ""


def test_lookup_prefix():
    mp = {"T25.300": "足部三度烧伤"}
    # 带扩展码，去掉 x 后前缀命中
    assert F.lookup("T25.300x005", mp) == "足部三度烧伤"


def test_fill_column_multivalue_aligned():
    codes = pd.Series(["I50.900|T31.000", "I10.x00x002", None])
    names = pd.Series([None, "高血压", None])
    mp = {"I50.900": "心力衰竭", "T31.000": "累及体表10%以下的烧伤",
          "I10.X00X002": "高血压"}
    out, filled = F.fill_column(codes, names, mp)
    # 仅 row0 被新填充；row1 已有名称，保留且不计入 filled
    assert filled == 1
    assert out[0] == "心力衰竭|累及体表10%以下的烧伤"
    assert out[1] == "高血压"  # 已有名称不覆盖
    assert out[2] == "" or pd.isna(out[2])


def test_build_maps_small(tmp_path):
    # 用两个极小的对照表验证 build_maps 收纳两列
    d10 = pd.DataFrame({
        "国临版编码": ["I50.000"], "国临版名称": ["心衰"],
        "医保版2.0编码": ["I50.900"], "医保版2.0名称": ["心力衰竭"],
    })
    d9 = pd.DataFrame({
        "国临3.0手术代码": ["47.0100"], "国临3.0手术名称": ["阑尾切除术"],
        "医保2.0手术代码": ["47.01"], "医保2.0手术名称": ["阑尾切除"],
    })
    d10.to_excel(tmp_path / "ICD10国临版2.0对照医保版2.0_0125.xlsx", index=False)
    d9.to_excel(tmp_path / "ICD9国临版3.0对照医保版2.0_0125.xlsx", index=False)
    diag_map, oprn_map = F.build_maps(str(tmp_path))
    assert diag_map.get("I50.900") == "心力衰竭"
    assert diag_map.get("I50.000") == "心衰"
    assert oprn_map.get("47.01") == "阑尾切除"
    assert oprn_map.get("47.0100") == "阑尾切除术"


def test_write_excel_textsafe_preserves_long_id(tmp_path):
    # 超长结算ID 必须以文本写出，不能被 Excel 存成科学计数法/丢精度
    df = pd.DataFrame({
        "record_id": [13018120260601001925402379, 13018120260608001934961252],
        "相关诊断编码": ["I50.900", "E11.900"],
    })
    out = tmp_path / "out.xlsx"
    F.write_excel_textsafe(df, str(out))
    back = pd.read_excel(out)
    assert back["record_id"].dtype == object  # 应为文本而非 float
    assert str(back["record_id"][0]) == "13018120260601001925402379"
    assert len(str(back["record_id"][0])) == 26
    # 单元格应为文本格式
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=2, column=1).number_format == "@"


def test_detect_text_id_columns_heuristic():
    df = pd.DataFrame({
        "record_id": ["13018120260601001925402379"],
        "主要诊断编码": ["I50.900"],   # 含字母，非纯数字 -> 不误判
        "医疗总费用": [1234.5],          # 数值 -> 不误判
        "年龄": [65],
    })
    cols = F._detect_text_id_columns(df)
    assert "record_id" in cols
    assert "主要诊断编码" not in cols
    assert "医疗总费用" not in cols
    assert "年龄" not in cols
